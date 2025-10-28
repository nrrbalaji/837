import openpyxl
import sys

try:
    wb = openpyxl.load_workbook(r'D:\AI\from my laptop\837\837_Segment__Matrix.xlsx')
    print(f"Sheet names: {wb.sheetnames}\n")
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"\n{'='*80}")
        print(f"SHEET: {sheet_name}")
        print(f"{'='*80}\n")
        
        max_rows = min(100, ws.max_row)  # Limit to first 100 rows
        
        for row in ws.iter_rows(min_row=1, max_row=max_rows, values_only=True):
            print('\t'.join([str(cell) if cell is not None else '' for cell in row]))
        
        if ws.max_row > 100:
            print(f"\n... (showing first 100 of {ws.max_row} rows)")
            
except Exception as e:
    print(f"Error: {e}", file=sys.stderr)
